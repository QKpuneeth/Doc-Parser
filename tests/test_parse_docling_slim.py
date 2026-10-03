import hashlib
import json
import re
import tempfile
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from PIL import Image, ImageDraw

import image_signals
import vision_utils
from parse_docling_slim import (
    IMAGE_MARKER,
    _apply_vision,
    _build_cli,
    _classify_images,
    _full_page_duplicate_signals,
    _image_metadata,
    _is_decorative,
    _merge_transcriptions,
    _refs_nested_in_pictures,
    _render_markdown,
    _semantic_loss_warnings,
    _table_representation,
    _validate_outputs,
    _vision_block,  # backwards-compatible alias of _image_block
    _image_block,
    parse_document,
    parse_many,
    ParserConfig,
)


def make_record(**overrides):
    record = {
        "type": "image",
        "ref": "#/pictures/0",
        "order": 0,
        "level": 1,
        "page": 1,
        "page_inferred": False,
        "source": {"page": 1, "bbox": [0, 0, 100, 100]},
        "relative_path": "images/image_0001.png",
        "path": "images/image_0001.png",
        "saved": True,
        "sha256": "a" * 64,
        "width": 100,
        "height": 100,
        "visual": {"width": 100, "height": 100, "area": 10000, "min_dimension": 100},
        "position": {"band": "middle-center", "left_ratio": 0.4, "top_ratio": 0.4},
        "coverage": 0.02,
        "background": False,
        "include_background": False,
        "include_decorative": False,
        "ocr": "",
        "include_in_md": True,
        "vision": None,
        "vision_status": None,
        "classification": None,
        "annotations": [],
        "chart_data": None,
    }
    record.update(overrides)
    return record


class DecorativeDetectionTests(unittest.TestCase):
    """A: logo. B: percentages. C: numbers. D: labelled diagram."""

    def decide(self, **overrides):
        records = overrides.pop("siblings", None)
        record = make_record(**overrides)
        group = list(records) if records else [record]
        group.append(record)
        _classify_images(group)
        return record

    def test_a_logo_is_decorative(self):
        record = self.decide(ocr="IQVIA", width=162, height=32)
        self.assertFalse(record["include_in_md"])
        self.assertEqual(record["omit_reason"], "decorative_or_logo")
        self.assertEqual(record["semantic_class"], "logo")
        self.assertTrue(_is_decorative(record))
        self.assertTrue(record["saved"])

    def test_a_garbled_logo_variant_is_decorative(self):
        record = self.decide(ocr="1IQVIA", width=163, height=33)
        self.assertFalse(record["include_in_md"])
        self.assertEqual(record["semantic_class"], "logo")

    def test_a_page_number_is_decorative(self):
        record = self.decide(ocr="12", width=40, height=20)
        self.assertFalse(record["include_in_md"])
        self.assertEqual(record["semantic_class"], "page_number")

    def test_b_small_chart_with_percentages_is_kept(self):
        record = self.decide(
            ocr="% of physicians\n14%\n19%\n35%\n28%",
            width=248,
            height=291,
            coverage=0.0361,
            relative_path="images/image_0042.png",
            path="images/image_0042.png",
        )
        self.assertTrue(record["include_in_md"])
        self.assertNotIn("omit_reason", record)
        self.assertEqual(record["semantic_class"], "chart")
        self.assertIn("contains percentage values", record["meaningful_evidence"])
        self.assertFalse(_is_decorative(record))

    def test_b_tilde_percentage_chart_is_kept(self):
        record = self.decide(ocr="~45%", width=90, height=60)
        self.assertTrue(record["include_in_md"])

    def test_c_small_chart_with_numbers_is_kept(self):
        record = self.decide(ocr="35\n28\n19\n14", width=120, height=80)
        self.assertTrue(record["include_in_md"])
        self.assertIn("contains multiple numeric values", record["meaningful_evidence"])

    def test_d_diagram_with_labels_is_kept(self):
        record = self.decide(
            ocr="Specialized centres\nEmergency Room",
            width=249,
            height=116,
        )
        self.assertTrue(record["include_in_md"])
        self.assertIn("contains meaningful text labels", record["meaningful_evidence"])

    def test_single_meaningful_word_is_not_decorative(self):
        record = self.decide(ocr="Others", width=162, height=34)
        self.assertTrue(record["include_in_md"])

    def test_size_alone_never_drops_an_image(self):
        for ocr in ("14%\n19%", "35\n28", "Adherence rate"):
            with self.subTest(ocr=ocr):
                record = self.decide(ocr=ocr, width=8, height=8, coverage=0.0001)
                self.assertTrue(record["include_in_md"])

    def test_repeated_page_furniture_without_text_is_dropped(self):
        logo_shape = {
            "width": 274,
            "height": 41,
            "area": 274 * 41,
            "min_dimension": 41,
            "aspect_ratio": round(274 / 41, 4),
            "thin_band": True,
            "tiny": False,
        }
        records = [
            make_record(
                page=3,
                position={"band": "top-right", "left_ratio": 0.82, "top_ratio": 0.13},
                visual=logo_shape,
                coverage=0.0056,
                ocr="",
            ),
            make_record(
                page=8,
                position={"band": "top-right", "left_ratio": 0.82, "top_ratio": 0.13},
                visual=logo_shape,
                coverage=0.004,
                ocr="",
            ),
        ]
        _classify_images(records)
        for record in records:
            self.assertFalse(record["include_in_md"])
            self.assertEqual(record["semantic_class"], "icon_or_decoration")

    def test_large_unclassified_image_is_preserved(self):
        record = self.decide(ocr="", width=208, height=302, coverage=0.031)
        self.assertTrue(record["include_in_md"])

    def test_include_decorative_flag_forces_markdown_inclusion(self):
        record = make_record(ocr="IQVIA", width=162, height=32, include_decorative=True)
        _classify_images([record])
        self.assertTrue(record["include_in_md"])

    def test_classification_and_evidence_are_recorded(self):
        record = self.decide(ocr="% of physicians\n14%\n35%", width=120, height=90)
        self.assertEqual(record["classification"], "chart")
        self.assertTrue(record["meaningful"])
        self.assertIn("signals", record)
        self.assertIn("has_percent", record["signals"]["text"])
        self.assertTrue(record["alt_text"])


class NestedPictureTextTests(unittest.TestCase):
    """Text inside a picture is only emitted by our own image block.

    ``export_to_markdown(traverse_pictures=False)`` drops nested text, so a
    page-sized picture must not be omitted as a duplicate of text the reader
    never sees. Regression: the first 45-page run silently deleted 2010 text
    blocks from document.md this way.
    """

    class _Node:
        def __init__(self, self_ref, children=()):
            self.self_ref = self_ref
            self.children = list(children)
            self.text = None

    def doc(self, nested_refs=(), top_refs=()):
        nested = [self._Node(ref) for ref in nested_refs]
        group = self._Node("#/groups/1", children=nested) if nested else None
        picture = self._Node("#/pictures/0", children=[group] if group else [])
        doc = type("_Doc", (), {})()
        doc.pictures = [picture]
        doc.standalone = [self._Node(ref) for ref in top_refs]
        return doc

    def test_collects_refs_of_items_inside_a_picture(self):
        doc = self.doc(nested_refs=("#/texts/1", "#/texts/2"))
        refs = _refs_nested_in_pictures(doc)
        # Intermediate group nodes are collected too, which is harmless.
        self.assertLessEqual({"#/texts/1", "#/texts/2"}, refs)

    def test_ignores_top_level_text(self):
        doc = self.doc(nested_refs=("#/texts/1",), top_refs=("#/texts/9",))
        refs = _refs_nested_in_pictures(doc)
        self.assertIn("#/texts/1", refs)
        self.assertNotIn("#/texts/9", refs)

    def test_picture_with_only_nested_text_is_never_a_duplicate(self):
        # The page text equals the picture OCR word for word, but that text is
        # nested inside the picture, so nothing else in the Markdown carries it.
        ocr = (
            "Distribution of TRD patient groups by phase. Phase 1 induction, "
            "Phase 2a maintenance, Phase 2b continuation, Phase 3 remission, "
            "Phase 4 recovery. Total n equals 57 respondents. Psychiatrists n 27, "
            "general practitioners n 30, patients with young adult onset."
        )
        record = {
            "relative_path": "images/image_0001.png",
            "page": 1,
            "coverage": 0.8,
            "ocr": ocr,
            "width": 1676,
            "height": 953,
        }
        config = ParserConfig()
        # Same words, but delivered as an exported page text layer.
        signals = _full_page_duplicate_signals(record, {1: [ocr]}, config)
        self.assertIsNotNone(signals)
        self.assertTrue(signals["duplicate"])

        # The same words nested in the picture: no page text is exported, so the
        # check does not apply and the picture is kept.
        self.assertIsNone(_full_page_duplicate_signals(record, {}, config))

    def test_resolves_docling_refitem_children_through_the_lookup(self):
        """Real Docling pictures hold RefItem children, not nodes.

        Modelling them as nodes made the first fix look correct while it did
        nothing: the guard never matched a single real document.
        """

        class RefItem:
            def __init__(self, cref):
                self.cref = cref

        text = self._Node("#/texts/1")
        group = self._Node("#/groups/1", children=[RefItem("#/texts/1")])
        picture = self._Node("#/pictures/0", children=[RefItem("#/groups/1")])
        doc = type("_Doc", (), {})()
        doc.pictures = [picture]
        lookup = {"#/groups/1": group, "#/texts/1": text}
        refs = _refs_nested_in_pictures(doc, lookup)
        self.assertIn("#/texts/1", refs)
        self.assertIn("#/groups/1", refs)

    def test_nested_text_does_not_enter_the_comparison_text(self):
        # The guard is applied where page text is gathered, so a nested item
        # must be excluded from the layer the comparison sees.
        doc = self.doc(nested_refs=("#/texts/1",), top_refs=("#/texts/2",))
        nested = _refs_nested_in_pictures(doc)
        layer = ["exported"]
        if "#/texts/1" not in nested:
            layer.append("nested")
        self.assertEqual(layer, ["exported"])


class FullPageDuplicateDetectionTests(unittest.TestCase):
    """CHANGE 4: a page-sized image that only repeats the page text layer is
    redundant, but a figure with unique content must always survive."""

    PAGE_TEXT = (
        "Spravato is indicated for the treatment of treatment-resistant depression in "
        "adults. Induction dosing should be administered twice weekly for the first "
        "four weeks under direct supervision in a certified healthcare setting with at "
        "least three hours of monitoring. 45 percent of psychiatrists and 28 percent of "
        "general practitioners recommend specialist centres for the induction phase, "
        "while 19 percent prefer outpatient administration. Maintenance treatment may "
        "continue in an outpatient setting following successful induction. Patients must "
        "be monitored for sedation, dissociation, and blood pressure changes. Dose "
        "adjustments are made weekly based on response and tolerability over a "
        "maintenance period of six months or longer."
    )

    def page_image(self, **overrides):
        defaults = {
            "coverage": 0.92,
            "background": True,
            "include_background": False,
            "width": 1200,
            "height": 1600,
            "visual": {
                "width": 1200,
                "height": 1600,
                "area": 1920000,
                "min_dimension": 1200,
                "flat": False,
                "near_blank": False,
                "tiny": False,
                "thin_band": False,
            },
        }
        defaults.update(overrides)
        return make_record(**defaults)

    def decide(self, record, page_text, config=None):
        _classify_images([record], config=config, page_text=page_text)
        return record

    def test_page_image_repeating_page_text_is_omitted(self):
        record = self.decide(
            self.page_image(ocr=self.PAGE_TEXT), {1: [self.PAGE_TEXT, "Source: IQVIA."]}
        )
        self.assertFalse(record["include_in_md"])
        self.assertEqual(record["semantic_class"], "duplicate")
        self.assertEqual(record["omit_reason"], "full_page_image_duplicates_text_layer")
        self.assertTrue(record["text_layer_duplicate"]["duplicate"])
        self.assertEqual(record["text_layer_duplicate"]["token_coverage"], 1.0)
        # The asset itself is still extracted: only its Markdown entry is dropped.
        self.assertTrue(record["saved"])

    def test_page_image_with_unique_content_is_kept(self):
        unique = (
            "Protocol step chart: screen patients, confirm diagnosis, administer 56 mg, "
            "observe for two hours, record response, schedule follow up. Response rate by "
            "site: North 61 percent, South 44 percent, East 52 percent, West 38 percent, "
            "Central 47 percent. Adverse events were reported in 12 percent of sessions "
            "and were transient in every case. Source: pivotal trial CS-200-009."
        )
        record = self.decide(self.page_image(ocr=unique), {1: [self.PAGE_TEXT]})
        self.assertTrue(record["include_in_md"])
        self.assertIsNone(record.get("omit_reason"))
        self.assertFalse(record["text_layer_duplicate"]["duplicate"])

    def test_similarity_threshold_governs_the_edge_case(self):
        """A page image that is more than the configured 0.90 similar is dropped.

        This is the documented trade-off of the default threshold: a short unique
        caption inside an otherwise duplicated page still scores above 0.90 and is
        therefore treated as redundant. Lowering ``full_page_duplicate_similarity``
        keeps such images, at the cost of retaining some genuinely duplicated ones.
        """
        mostly_duplicated = self.PAGE_TEXT + " An inset chart reports 73 percent concordance."
        record = self.decide(self.page_image(ocr=mostly_duplicated), {1: [self.PAGE_TEXT]})
        self.assertGreaterEqual(record["text_layer_duplicate"]["token_coverage"], 0.9)
        self.assertFalse(record["include_in_md"])

        from parse_docling_slim import ParserConfig

        kept = self.page_image(ocr=mostly_duplicated)
        self.decide(kept, {1: [self.PAGE_TEXT]}, config=ParserConfig(full_page_duplicate_similarity=0.99))
        self.assertTrue(kept["include_in_md"])

    def test_page_image_without_text_layer_is_kept(self):
        record = self.decide(self.page_image(ocr=self.PAGE_TEXT), {})
        self.assertTrue(record["include_in_md"])
        self.assertNotIn("text_layer_duplicate", record)

    def test_short_text_is_never_treated_as_duplicate(self):
        record = self.decide(
            self.page_image(ocr="45% of physicians recommend specialist centres."),
            {1: [self.PAGE_TEXT]},
        )
        self.assertTrue(record["include_in_md"])
        self.assertFalse(record["text_layer_duplicate"]["duplicate"])

    def test_ordinary_figure_is_not_compared(self):
        record = self.decide(
            self.page_image(ocr=self.PAGE_TEXT, coverage=0.1, background=False, width=400, height=300),
            {1: [self.PAGE_TEXT]},
        )
        self.assertTrue(record["include_in_md"])
        self.assertNotIn("text_layer_duplicate", record)

    def test_detection_can_be_disabled(self):
        from parse_docling_slim import ParserConfig

        record = self.decide(
            self.page_image(ocr=self.PAGE_TEXT),
            {1: [self.PAGE_TEXT]},
            config=ParserConfig(enable_full_page_duplicate_detection=False),
        )
        self.assertTrue(record["include_in_md"])
        self.assertNotIn("text_layer_duplicate", record)

    def test_similarity_primitive_is_token_presence_based(self):
        signals = image_signals.text_duplicate_signals(
            "Alpha beta gamma delta. 10 20 30", "Totally different words here. 10 20 30 40"
        )
        self.assertFalse(signals["duplicate"])
        self.assertLess(signals["token_coverage"], 0.9)

    def test_similarity_primitive_ignores_token_repetition(self):
        once = image_signals.text_duplicate_signals(
            self.PAGE_TEXT, self.PAGE_TEXT + " Extra trailing sentence."
        )
        thrice = image_signals.text_duplicate_signals(
            self.PAGE_TEXT + " " + self.PAGE_TEXT + " " + self.PAGE_TEXT, self.PAGE_TEXT
        )
        self.assertTrue(once["duplicate"])
        self.assertTrue(thrice["duplicate"])

    def test_text_layer_duplicate_is_not_reported_as_semantic_loss(self):
        import tempfile as _tempfile

        record = self.decide(
            self.page_image(ocr=self.PAGE_TEXT), {1: [self.PAGE_TEXT, "Source: IQVIA."]}
        )
        with _tempfile.TemporaryDirectory() as tmp:
            findings = _semantic_loss_warnings([record], Path(tmp))
        self.assertEqual(findings, [])


class MarkdownImagePolicyTests(unittest.TestCase):
    def test_image_block_contains_ocr_and_vision(self):
        record = {
            "include_in_md": True,
            "relative_path": "images/image_0001.png",
            "ocr": "Visible title",
            "vision": {
                "status": "ok",
                "category": "diagram_chart",
                "description": "A chart.",
                "extracted_text": "Chart label",
            },
        }
        rendered = _render_markdown(f"before {IMAGE_MARKER} after", [record])
        self.assertIn("![Visible title](images/image_0001.png)", rendered)
        self.assertIn("**Visual description:**", rendered)
        self.assertIn("A chart.", rendered)
        self.assertIn("**OCR text:**", rendered)
        self.assertIn("Visible title", rendered)
        self.assertNotIn(IMAGE_MARKER, rendered)

    def test_ocr_and_vision_are_never_labelled_as_each_other(self):
        record = make_record(
            ocr="Native OCR line",
            semantic_class="chart",
            include_in_md=True,
            vision={"status": "ok", "category": "diagram_chart", "description": "Vision line.",
                    "extracted_text": "Vision transcription"},
        )
        block = _image_block(record)
        ocr_section = block.split("**OCR text:**", 1)[1].split("**Visual description:**", 1)[0]
        vision_section = block.split("**Visual description:**", 1)[1]
        self.assertIn("Native OCR line", ocr_section)
        self.assertNotIn("Vision line.", ocr_section)
        self.assertIn("Vision line.", vision_section)

    def test_missing_ocr_or_vision_omits_the_section(self):
        record = make_record(ocr="", semantic_class="chart", include_in_md=True)
        block = _image_block(record)
        self.assertNotIn("**OCR text:**", block)
        self.assertNotIn("**Visual description:**", block)
        self.assertIn("**Image type:** chart", block)

        record = make_record(ocr="Only native text", semantic_class="diagram", include_in_md=True)
        block = _image_block(record)
        self.assertIn("**OCR text:**", block)
        self.assertNotIn("**Visual description:**", block)
        self.assertNotIn("not available", block.lower())
        self.assertNotIn("unavailable", block.lower())

    def test_image_path_is_relative_and_portable(self):
        record = make_record(ocr="% of physicians\n14%", semantic_class="chart", include_in_md=True)
        block = _image_block(record)
        self.assertIn("(images/image_0001.png)", block)
        self.assertNotIn("/home/", block)
        self.assertNotIn("C:\\", block)
        self.assertNotIn("data:", block)

    def test_image_ocr_can_be_disabled(self):
        record = {
            "include_in_md": True,
            "relative_path": "images/image_0001.png",
            "ocr": "Visible title",
        }
        rendered = _render_markdown(IMAGE_MARKER, [record], include_ocr=False)
        self.assertIn("images/image_0001.png", rendered)
        self.assertNotIn("Visible title", rendered)

    def test_decorative_image_is_traceable_without_link(self):
        record = {
            "include_in_md": False,
            "omit_reason": "decorative_or_logo",
            "relative_path": "images/image_0001.png",
        }
        rendered = _render_markdown(IMAGE_MARKER, [record])
        self.assertIn("image omitted: decorative_or_logo", rendered)
        self.assertIn("images/image_0001.png", rendered)
        self.assertNotIn("![", rendered)

    def test_marker_mismatch_fails(self):
        with self.assertRaises(RuntimeError):
            _render_markdown(IMAGE_MARKER, [])

    def test_no_base64_payloads_are_emitted(self):
        record = make_record(
            ocr="% of physicians\n14%",
            semantic_class="chart",
            include_in_md=True,
        )
        _classify_images([record])
        block = _image_block(record)
        self.assertIn("![", block)
        self.assertIn("(images/", block)
        self.assertNotIn("base64", block)
        self.assertNotIn("data:image", block)

    def test_image_without_semantics_stays_in_markdown(self):
        record = make_record(ocr="", semantic_class="unclassified_image", include_in_md=True)
        _classify_images([record])
        block = _image_block(record)
        self.assertIn("images/", block)
        self.assertIn("![", block)
        self.assertNotIn("Not available", block)

    def test_percentage_chart_block_keeps_ocr_values(self):
        record = make_record(
            ocr="% of physicians\n14%\n19%\n35%\n28%",
            page=32,
            width=248,
            height=291,
        )
        _classify_images([record])
        block = _image_block(record)
        self.assertIn("![% of physicians chart](images/image_0001.png)", block)
        for value in ("14%", "19%", "35%", "28%"):
            self.assertIn(value, block)
        self.assertIn("**Image type:** chart", block)
        self.assertIn("**OCR text:**", block)


class SemanticTextMergeTests(unittest.TestCase):
    """OCR and Vision transcription share one de-duplicated Markdown section."""

    def test_vision_disabled_keeps_ocr_only_behaviour(self):
        record = make_record(
            ocr="Native OCR line",
            semantic_class="chart",
            vision=None,
            vision_status="disabled",
        )
        block = _image_block(record)
        metadata = _image_metadata(record)
        self.assertIn("**OCR text:**", block)
        self.assertIn("Native OCR line", block)
        self.assertNotIn("**Visual description:**", block)
        self.assertNotIn("**Vision-transcribed text:**", block)
        self.assertEqual(metadata["text_source"], "native_ocr")
        self.assertEqual(metadata["semantic_text"], "Native OCR line")

    def test_identical_ocr_and_vision_transcription_deduplicates(self):
        record = make_record(
            ocr="57 HCPs were surveyed in phase 3\nOntario 47%",
            semantic_class="chart",
            vision={
                "status": "ok",
                "description": "A chart.",
                "extracted_text": "57 HCPs were surveyed in phase 3\nOntario 47%",
            },
        )
        block = _image_block(record)
        self.assertIn("**OCR text:**", block)
        self.assertIn("**Visual description:**", block)
        self.assertNotIn("**Vision-transcribed text:**", block)
        ocr_body = block.split("**OCR text:**", 1)[1].split("**Visual description:**", 1)[0]
        self.assertEqual(ocr_body.count("57 HCPs were surveyed"), 1)
        metadata = _image_metadata(record)
        self.assertEqual(metadata["text_source"], "native_ocr")
        # The rich Vision transcription is preserved internally for audit.
        self.assertEqual(metadata["vision"]["extracted_text"], "57 HCPs were surveyed in phase 3\nOntario 47%")

    def test_formatting_differences_are_still_a_duplicate(self):
        record = make_record(
            ocr="57 HCPs were surveyed in phase 3.",
            semantic_class="chart",
            vision={"status": "ok", "description": "A chart.", "extracted_text": "57 HCPs were surveyed in phase 3"},
        )
        block = _image_block(record)
        self.assertIn("**OCR text:**", block)
        self.assertNotIn("**Vision-transcribed text:**", block)
        metadata = _image_metadata(record)
        self.assertEqual(metadata["text_source"], "native_ocr")
        self.assertEqual(metadata["semantic_text"], "57 HCPs were surveyed in phase 3.")

    def test_vision_text_with_new_information_is_merged(self):
        record = make_record(
            ocr="Ontario 47%",
            semantic_class="chart",
            vision={
                "status": "ok",
                "description": "A chart.",
                "extracted_text": "Ontario 47%\nQuebec 23%\nWest 21%\nEast 9%",
            },
        )
        block = _image_block(record)
        self.assertNotIn("**Vision-transcribed text:**", block)
        ocr_section = block.split("**OCR text:**", 1)[1].split("**Visual description:**", 1)[0]
        for value in ("Ontario 47%", "Quebec 23%", "West 21%", "East 9%"):
            self.assertIn(value, ocr_section)
        self.assertEqual(ocr_section.count("Ontario"), 1)  # common text is not duplicated
        metadata = _image_metadata(record)
        self.assertEqual(metadata["text_source"], "merged")
        self.assertIn("Quebec 23%", metadata["semantic_text"])

    def test_vision_text_is_used_when_native_ocr_is_empty(self):
        record = make_record(
            ocr="",
            semantic_class="diagram",
            vision={"status": "ok", "description": "A diagram.", "extracted_text": "refined approach\n10 steps"},
        )
        block = _image_block(record)
        self.assertIn("**OCR text:**", block)
        self.assertIn("refined approach", block)
        self.assertNotIn("**Vision-transcribed text:**", block)
        metadata = _image_metadata(record)
        self.assertEqual(metadata["text_source"], "vision")
        self.assertEqual(metadata["vision"]["extracted_text"], "refined approach\n10 steps")

    def test_description_only_image_keeps_type_and_description(self):
        record = make_record(
            ocr="",
            semantic_class="chart",
            vision={"status": "ok", "description": "Regional demographics.", "extracted_text": ""},
        )
        block = _image_block(record)
        self.assertIn("![", block)
        self.assertIn("**Image type:** chart", block)
        self.assertIn("**Visual description:**", block)
        self.assertIn("Regional demographics.", block)
        self.assertNotIn("**OCR text:**", block)
        self.assertNotIn("**Vision-transcribed text:**", block)
        self.assertNotIn("not available", block.lower())

    def test_no_content_keeps_reference_without_placeholders(self):
        record = make_record(
            ocr="",
            semantic_class="chart",
            vision={"status": "ok", "description": "", "extracted_text": ""},
        )
        block = _image_block(record)
        self.assertIn("![", block)
        self.assertIn("**Image type:** chart", block)
        self.assertNotIn("**OCR text:**", block)
        self.assertNotIn("**Visual description:**", block)
        self.assertNotIn("not available", block.lower())
        self.assertNotIn("unavailable", block.lower())
        metadata = _image_metadata(record)
        self.assertEqual(metadata["text_source"], "none")
        self.assertEqual(metadata["semantic_text"], "")

    def test_vision_api_failure_leaves_ocr_as_the_only_transcription(self):
        record = make_record(
            ocr="Native OCR line",
            semantic_class="chart",
            vision={"status": "error", "error_type": "APIConnectionError", "description": "", "extracted_text": ""},
            vision_status="error",
            vision_error="APIConnectionError",
        )
        block = _image_block(record)
        self.assertIn("**OCR text:**", block)
        self.assertIn("Native OCR line", block)
        self.assertNotIn("**Visual description:**", block)
        self.assertNotIn("**Vision-transcribed text:**", block)
        metadata = _image_metadata(record)
        self.assertEqual(metadata["text_source"], "native_ocr")

    def test_merge_primitive_distinguishes_source_provenance(self):
        self.assertEqual(_merge_transcriptions("Ontario 47%", "Ontario 47%")["source"], "native_ocr")
        self.assertEqual(_merge_transcriptions("", "Ontario 47%")["source"], "vision")
        self.assertEqual(_merge_transcriptions("Ontario 47%", "")["source"], "native_ocr")
        self.assertEqual(_merge_transcriptions("", "")["source"], "none")
        merged = _merge_transcriptions("Ontario 47%", "Ontario 47% Quebec 23%")
        self.assertEqual(merged["source"], "merged")
        self.assertIn("Quebec 23%", merged["text"])

    def test_vision_transcription_is_never_renderable_as_a_separate_section(self):
        record = make_record(
            ocr="Real OCR",
            semantic_class="chart",
            vision={"status": "ok", "description": "Desc.", "extracted_text": "Anything at all"},
        )
        rendered = _render_markdown(f"before {IMAGE_MARKER} after", [record])
        self.assertNotIn("**Vision-transcribed text:**", rendered)
        self.assertIn("**OCR text:**", rendered)
        self.assertIn("**Visual description:**", rendered)


class VisionUtilsTests(unittest.TestCase):
    def test_structured_response_and_fenced_json(self):
        response = SimpleNamespace(
            choices=[
                SimpleNamespace(
                    message=SimpleNamespace(
                        content='```json\n{"category":"diagram_chart","extracted_text":"42%","description":"A chart."}\n```'
                    )
                )
            ],
            usage=SimpleNamespace(prompt_tokens=2, completion_tokens=3, total_tokens=5),
        )
        with tempfile.TemporaryDirectory() as directory:
            image_path = Path(directory) / "image.png"
            Image.new("RGB", (10, 10), "white").save(image_path)
            with patch.object(vision_utils, "VISION_ENABLED", True), patch.object(
                vision_utils, "_IS_AZURE", False
            ), patch.object(vision_utils, "API_KEY", "test"), patch.object(
                vision_utils, "API_BASE", "https://example.invalid"
            ), patch.object(vision_utils, "API_VERSION", ""), patch.object(
                vision_utils, "VISION_PROVIDER", "openai"
            ), patch.object(
                vision_utils, "completion", return_value=response
            ) as completion, patch.object(
                vision_utils, "completion_cost", return_value=0.01
            ):
                result = vision_utils.classify_and_extract_image(str(image_path), timeout=1)
        self.assertEqual(result["category"], "diagram_chart")
        self.assertEqual(result["extracted_text"], "42%")
        self.assertEqual(result["cost_usd"], 0.01)
        self.assertEqual(completion.call_args.kwargs["model"], "gpt-4o-mini")


class VisionDecisionTests(unittest.TestCase):
    """F, G, H: bad OCR, Vision disabled, Vision failure."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.output = Path(self.temp_dir.name)
        (self.output / "images").mkdir(parents=True, exist_ok=True)
        for name in ("image_0001.png", "image_0002.png", "image_0003.png"):
            Image.new("RGB", (40, 40), "white").save(self.output / "images" / name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def config(self, **overrides):
        from parse_docling_slim import ParserConfig

        values = {"enable_vision": True, "vision_timeout": 1.0}
        values.update(overrides)
        return ParserConfig(**values)

    def test_f_garbled_ocr_image_is_kept(self):
        record = make_record(ocr="\ufffd\ufffd 1 2\n| | \ufffd", width=300, height=200)
        _classify_images([record])
        self.assertTrue(record["include_in_md"])
        self.assertTrue(record["signals"]["text"]["garbled"])

    def test_g_vision_disabled_succeeds_and_is_recorded(self):
        records = [make_record(ocr="14%\n35%"), make_record(ocr="IQVIA", width=162, height=32)]
        _classify_images(records)
        stats = _apply_vision(records, self.output, self.config(enable_vision=False))
        self.assertEqual(stats["calls"], 0)
        self.assertTrue(all(record["vision_status"] == "disabled" for record in records))
        self.assertEqual(stats["statuses"]["disabled"], 2)

    def test_h_vision_failure_does_not_fail_the_run(self):
        # An image whose classification is uncertain: that is exactly the case
        # CHANGE 3 sends to Vision, so the failure path is exercised.
        record = make_record(ocr="\ufffd\ufffd 1 2\n| | \ufffd", width=300, height=200)
        _classify_images([record])
        self.assertLess(record["classification_confidence"], 0.75)
        with patch(
            "parse_docling_slim.classify_and_extract_image",
            side_effect=RuntimeError("boom"),
        ):
            stats = _apply_vision([record], self.output, self.config())
        self.assertEqual(stats["errors"], 1)
        self.assertEqual(record["vision_status"], "error")
        self.assertEqual(record["vision_error"], "RuntimeError")
        self.assertEqual(stats["retried"], 0)
        block = _vision_block(record)
        self.assertIn("images/image_0001.png", block)

    def test_vision_prioritises_charts_and_skips_decorative(self):
        chart = make_record(ocr="14%\n19%\n35%\n28%", order=1, semantic_class="chart")
        logo = make_record(ocr="IQVIA", order=2, width=162, height=32)
        plain = make_record(ocr="Some caption text", order=3)
        records = [chart, logo, plain]
        _classify_images(records)
        seen = []

        def fake(path, enabled=True, timeout=0):
            seen.append(Path(path).name)
            return {
                "category": "diagram_chart",
                "extracted_text": "42%",
                "description": "A pie chart.",
                "status": "ok",
                "cost_usd": 0.002,
                "usage_tokens": {"input_tokens": 10, "output_tokens": 5, "total_tokens": 15},
            }

        with patch("parse_docling_slim.vision_available", return_value=True), patch(
            "parse_docling_slim.classify_and_extract_image", side_effect=fake
        ):
            # Confidence gate open: no image reaches a confidence of 1.0, so this
            # test isolates priority ordering rather than the CHANGE 3 gate.
            stats = _apply_vision(records, self.output, self.config(vision_confidence_threshold=1.0))
        self.assertEqual(stats["calls"], 2)
        self.assertEqual(seen[0], "image_0001.png")
        self.assertEqual(len(seen), 2)
        self.assertEqual(logo["vision_status"], "skipped_decorative")
        self.assertGreater(stats["tokens"]["total"], 0)
        self.assertAlmostEqual(stats["cost_usd"], 0.004)

    def test_vision_max_images_limit_is_reported(self):
        records = [make_record(ocr=f"{value}%", order=index) for index, value in enumerate((14, 19, 35))]
        _classify_images(records)
        with patch("parse_docling_slim.vision_available", return_value=True), patch(
            "parse_docling_slim.classify_and_extract_image",
            return_value={"category": "other", "extracted_text": "", "description": "", "status": "ok"},
        ):
            stats = _apply_vision(
                records, self.output, self.config(vision_max_images=1, vision_confidence_threshold=1.0)
            )
        self.assertEqual(stats["calls"], 1)
        self.assertEqual(stats["statuses"].get("skipped_limit"), 2)

    def test_negative_vision_max_images_is_rejected(self):
        records = [make_record()]
        with self.assertRaises(ValueError):
            _apply_vision(records, self.output, self.config(vision_max_images=-1))


class VisionConfidenceGateTests(unittest.TestCase):
    """CHANGE 3: Vision is spent on uncertain images, retried when transient."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.output = Path(self.temp_dir.name)
        (self.output / "images").mkdir(parents=True, exist_ok=True)
        for name in ("image_0001.png", "image_0002.png"):
            Image.new("RGB", (40, 40), "white").save(self.output / "images" / name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def config(self, **overrides):
        from parse_docling_slim import ParserConfig

        values = {"enable_vision": True, "vision_timeout": 1.0}
        values.update(overrides)
        return ParserConfig(**values)

    def ok_result(self, path, enabled=True, timeout=0):
        return {
            "category": "diagram_chart",
            "extracted_text": "42%",
            "description": "A pie chart.",
            "status": "ok",
            "cost_usd": 0.001,
            "usage_tokens": {"input_tokens": 10, "output_tokens": 5, "total_tokens": 15},
        }

    def test_confident_image_is_not_sent_to_vision(self):
        record = make_record(ocr="14%\n19%\n35%\n28%", order=1)
        _classify_images([record])
        self.assertGreaterEqual(record["classification_confidence"], 0.75)
        with patch("parse_docling_slim.vision_available", return_value=True), patch(
            "parse_docling_slim.classify_and_extract_image"
        ) as call:
            stats = _apply_vision([record], self.output, self.config())
        call.assert_not_called()
        self.assertEqual(stats["calls"], 0)
        self.assertEqual(record["vision_status"], "skipped_confident")
        self.assertIn("threshold", record["vision_skip_reason"])
        self.assertEqual(stats["statuses"]["skipped_confident"], 1)
        self.assertEqual(stats["confidence_threshold"], 0.75)

    def test_uncertain_image_is_sent_to_vision(self):
        record = make_record(ocr="\ufffd\ufffd 1 2\n| | \ufffd", width=300, height=200, order=1)
        _classify_images([record])
        with patch("parse_docling_slim.vision_available", return_value=True), patch(
            "parse_docling_slim.classify_and_extract_image", side_effect=self.ok_result
        ):
            stats = _apply_vision([record], self.output, self.config())
        self.assertEqual(stats["calls"], 1)
        self.assertEqual(record["vision_status"], "ok")
        self.assertEqual(record["vision_attempts"], 1)

    def test_lower_threshold_enables_more_calls(self):
        record = make_record(ocr="14%\n19%\n35%\n28%", order=1)
        _classify_images([record])
        with patch("parse_docling_slim.vision_available", return_value=True), patch(
            "parse_docling_slim.classify_and_extract_image", side_effect=self.ok_result
        ):
            stats = _apply_vision([record], self.output, self.config(vision_confidence_threshold=0.99))
        self.assertEqual(stats["calls"], 1)

    def test_required_image_bypasses_the_gate(self):
        record = make_record(ocr="14%\n19%\n35%\n28%", order=1)
        _classify_images([record])
        record["vision_required"] = True
        record["vision_required_reason"] = "figure referenced by surrounding text"
        with patch("parse_docling_slim.vision_available", return_value=True), patch(
            "parse_docling_slim.classify_and_extract_image", side_effect=self.ok_result
        ):
            stats = _apply_vision([record], self.output, self.config())
        self.assertEqual(stats["calls"], 1)
        self.assertEqual(record["vision_status"], "ok")

    def test_transient_failure_is_retried(self):
        record = make_record(ocr="\ufffd\ufffd 1 2\n| | \ufffd", width=300, height=200, order=1)
        _classify_images([record])
        responses = [
            {"status": "error", "error_type": "APITimeoutError"},
            self.ok_result("x"),
        ]
        with patch("parse_docling_slim.vision_available", return_value=True), patch(
            "parse_docling_slim.classify_and_extract_image", side_effect=responses
        ), patch("parse_docling_slim.time.sleep"):
            stats = _apply_vision([record], self.output, self.config(vision_retry_count=2))
        self.assertEqual(stats["retried"], 1)
        self.assertEqual(stats["api_calls"], 2)
        self.assertEqual(record["vision_attempts"], 2)
        self.assertEqual(record["vision_status"], "ok")

    def test_permanent_failure_is_not_retried(self):
        record = make_record(ocr="\ufffd\ufffd 1 2\n| | \ufffd", width=300, height=200, order=1)
        _classify_images([record])
        with patch("parse_docling_slim.vision_available", return_value=True), patch(
            "parse_docling_slim.classify_and_extract_image",
            return_value={"status": "error", "error_type": "BadRequestError"},
        ):
            stats = _apply_vision([record], self.output, self.config(vision_retry_count=3))
        self.assertEqual(stats["retried"], 0)
        self.assertEqual(stats["api_calls"], 1)
        self.assertEqual(stats["errors"], 1)
        self.assertEqual(record["vision_error"], "BadRequestError")
        self.assertEqual(record["vision_attempts"], 1)

    def test_raised_exception_is_recorded_not_fatal(self):
        record = make_record(ocr="\ufffd\ufffd 1 2\n| | \ufffd", width=300, height=200, order=1)
        _classify_images([record])
        with patch("parse_docling_slim.vision_available", return_value=True), patch(
            "parse_docling_slim.classify_and_extract_image", side_effect=TimeoutError("slow")
        ):
            stats = _apply_vision([record], self.output, self.config(vision_retry_count=0))
        self.assertEqual(stats["errors"], 1)
        self.assertEqual(record["vision_error"], "TimeoutError")

    def test_limit_reach_is_observable(self):
        records = [make_record(ocr=f"{value}%", order=index) for index, value in enumerate((14, 19, 35))]
        _classify_images(records)
        with patch("parse_docling_slim.vision_available", return_value=True), patch(
            "parse_docling_slim.classify_and_extract_image", side_effect=self.ok_result
        ):
            stats = _apply_vision(
                records,
                self.output,
                self.config(vision_max_images=1, vision_confidence_threshold=1.0),
            )
        self.assertTrue(stats["limit_reached"])
        self.assertIsNotNone(stats["limit_skipped_highest_priority"])
        skipped = [r for r in records if r.get("vision_status") == "skipped_limit"]
        self.assertTrue(skipped)
        for record in skipped:
            self.assertIn("higher priority", record["vision_skip_reason"])

    def test_transient_error_classification(self):
        from parse_docling_slim import _vision_error_is_transient

        for name in ("APITimeoutError", "RateLimitError", "ConnectionError", "ReadTimeout"):
            self.assertTrue(_vision_error_is_transient(name), name)
        for name in ("BadRequestError", "ValueError", "", None):
            self.assertFalse(_vision_error_is_transient(name), name)


class AssetDeduplicationTests(unittest.TestCase):
    """CHANGE 5: one stored asset, many occurrences, one consistent decision."""

    def logo(self, order, page, **overrides):
        return make_record(
            order=order,
            page=page,
            ocr="IQVIA",
            width=162,
            height=32,
            sha256="b" * 64,
            relative_path="images/image_0001.png",
            path="images/image_0001.png",
            **overrides,
        )

    def test_repeated_asset_gets_one_owner_and_occurrences(self):
        from parse_docling_slim import _link_duplicate_assets

        first = self.logo(order=0, page=1)
        second = self.logo(order=5, page=12, ref="#/pictures/9")
        chart = make_record(order=6, page=13, ocr="14%\n19%\n35%", sha256="c" * 64)
        summary = _link_duplicate_assets([first, second, chart])
        self.assertEqual(summary["unique_assets"], 2)
        self.assertEqual(summary["duplicate_assets"], 1)
        self.assertEqual(summary["duplicate_occurrences"], 1)
        self.assertIsNone(first["duplicate_of"])
        self.assertEqual(second["duplicate_of"], "images/image_0001.png")
        self.assertEqual(first["asset_hash"], "b" * 64)
        self.assertEqual(chart["duplicate_of"], None)
        self.assertEqual(chart["duplicate_occurrences"], [])
        self.assertEqual([o["page"] for o in first["duplicate_occurrences"]], [1, 12])
        self.assertTrue(first["duplicate_occurrences"][0]["is_owner"])
        self.assertFalse(second["duplicate_occurrences"][1]["is_owner"])
        self.assertEqual(summary["pages_per_asset"]["b" * 64], [1, 12])

    def test_decision_is_propagated_to_every_occurrence(self):
        from parse_docling_slim import _link_duplicate_assets

        first = self.logo(order=0, page=1)
        # A later occurrence of the same asset could be classified differently by
        # naive per-occurrence logic; the owner decision must win.
        second = self.logo(
            order=5, page=12, ref="#/pictures/9", semantic_class="chart", include_in_md=True
        )
        _classify_images([first])
        first["include_in_md"] = False
        first["omit_reason"] = "decorative_or_logo"
        first["semantic_class"] = "logo"
        first["classification_confidence"] = 0.87
        _link_duplicate_assets([first, second])
        self.assertEqual(second["semantic_class"], "logo")
        self.assertFalse(second["include_in_md"])
        self.assertEqual(second["omit_reason"], "decorative_or_logo")
        self.assertEqual(second["classification_confidence"], 0.87)

    def test_repeated_chart_is_not_omitted(self):
        from parse_docling_slim import _link_duplicate_assets

        records = [
            make_record(order=index, page=index + 1, ocr="14%\n19%\n35%", sha256="d" * 64)
            for index in range(3)
        ]
        _classify_images(records)
        summary = _link_duplicate_assets(records)
        self.assertEqual(summary["duplicate_occurrences"], 2)
        self.assertTrue(all(record["include_in_md"] for record in records))
        self.assertEqual(records[1]["duplicate_of"], "images/image_0001.png")

    def test_single_occurrence_is_untouched(self):
        from parse_docling_slim import _link_duplicate_assets

        record = self.logo(order=0, page=1)
        _classify_images([record])
        summary = _link_duplicate_assets([record])
        self.assertEqual(summary["unique_assets"], 1)
        self.assertIsNone(record["duplicate_of"])
        self.assertEqual(record["duplicate_occurrences"], [])
        self.assertEqual(record["semantic_class"], "logo")


class _ImageDocumentFixture:
    """A real two page PDF with a bar chart and a branding logo.

    Docling's layout model only reports a region as a picture when it looks like
    one, so the chart is drawn as an actual chart (axes, bars, labels) rather
    than a white rectangle with a caption. Parsing a real document keeps these
    tests honest: the artifacts under test only exist once a picture has really
    been extracted.
    """

    def __init__(self, root):
        self.root = Path(root)
        self.source = self.root / "sample.pdf"
        chart = self._bar_chart()
        logo = self._logo()
        self._write_pdf(chart, logo)

    def _bar_chart(self):
        from PIL import Image, ImageDraw

        image = Image.new("RGB", (900, 600), "white")
        draw = ImageDraw.Draw(image)
        draw.rectangle([80, 60, 820, 520], outline="black", width=3)
        for index, (height, label) in enumerate([(320, "14%"), (200, "19%"), (120, "35%"), (80, "28%")]):
            left = 140 + index * 170
            draw.rectangle([left, 500 - height, left + 90, 500], fill=(40, 90, 160))
            draw.text((left + 20, 505), label, fill="black")
        draw.text((330, 25), "% of physicians", fill="black")
        path = self.root / "chart.png"
        image.save(path)
        return path

    def _logo(self):
        from PIL import Image, ImageDraw

        image = Image.new("RGB", (320, 64), "white")
        draw = ImageDraw.Draw(image)
        draw.rectangle([4, 12, 60, 52], fill=(200, 30, 30))
        draw.text((80, 26), "IQVIA", fill="black")
        path = self.root / "logo.png"
        image.save(path)
        return path

    def _write_pdf(self, chart, logo):
        import pymupdf

        document = pymupdf.open()
        for number in (1, 2):
            page = document.new_page(width=595, height=842)
            page.insert_text((60, 80), f"Page {number}", fontsize=14)
            page.insert_text(
                (60, 110),
                "Induction dosing is administered twice weekly for the first four weeks.",
                fontsize=11,
            )
            page.insert_text(
                (60, 130),
                "Maintenance treatment may continue in an outpatient setting.",
                fontsize=11,
            )
            page.insert_image(pymupdf.Rect(60, 300, 510, 600), filename=str(chart))
            page.insert_image(pymupdf.Rect(60, 40, 180, 76), filename=str(logo))
        document.save(str(self.source))
        document.close()


class MetadataAndDecisionLogTests(unittest.TestCase):
    """CHANGE 7/8/9/14: optional canonical, metadata index, decision log."""

    @classmethod
    def setUpClass(cls):
        cls._temp = tempfile.TemporaryDirectory()
        cls.fixture = _ImageDocumentFixture(cls._temp.name)
        cls.default_output = Path(cls._temp.name) / "default"
        cls.result = parse_document(
            str(cls.fixture.source), str(cls.default_output), force=True, enable_vision=False
        )

    @classmethod
    def tearDownClass(cls):
        cls._temp.cleanup()

    def parse_with(self, name, **overrides):
        output = Path(self._temp.name) / name
        values = {"force": True, "enable_vision": False}
        values.update(overrides)
        result = parse_document(str(self.fixture.source), str(output), **values)
        return result, output

    def read(self, output, name):
        return json.loads((output / name).read_text(encoding="utf-8"))

    def names(self, output):
        return sorted(path.name for path in output.iterdir())

    def test_fixture_really_contains_images(self):
        self.assertGreaterEqual(self.result["stats"]["images_total"], 2)
        self.assertEqual(self.result["stats"]["pages"], 2)

    def test_default_output_has_markdown_metadata_manifest_only(self):
        output = self.default_output
        self.assertEqual(
            self.names(output), ["document.md", "images", "logs", "manifest.json", "metadata.json"]
        )
        manifest = self.read(output, "manifest.json")
        self.assertIsNone(manifest["emitted"]["canonical"])
        self.assertEqual(manifest["emitted"]["metadata"], "metadata.json")
        listed = {item["path"] for item in manifest["artifacts"]}
        self.assertIn("metadata.json", listed)
        self.assertIn("logs/image_decisions.jsonl", listed)
        self.assertNotIn("canonical.json", listed)
        # The in-memory contract of parse_document is unchanged.
        self.assertIn("images", self.result)
        self.assertIn("validation", self.result)

    def test_emit_canonical_adds_the_file_everywhere(self):
        result, output = self.parse_with("canonical", emit_canonical=True)
        self.assertIn("canonical.json", self.names(output))
        manifest = self.read(output, "manifest.json")
        self.assertIn("canonical.json", {item["path"] for item in manifest["artifacts"]})
        self.assertEqual(self.read(output, "metadata.json")["artifacts"]["canonical"], "canonical.json")
        self.assertEqual(result["validation"]["status"], "pass")

    def test_metadata_can_be_disabled(self):
        _, output = self.parse_with("nometa", emit_metadata=False)
        self.assertNotIn("metadata.json", self.names(output))
        self.assertIsNone(self.read(output, "manifest.json")["emitted"]["metadata"])

    def test_metadata_describes_the_document_and_every_image(self):
        from parse_docling_slim import SCHEMA_VERSION

        output = self.default_output
        metadata = self.read(output, "metadata.json")
        self.assertEqual(metadata["schema_version"], SCHEMA_VERSION)
        self.assertEqual(metadata["document"]["page_count"], self.result["stats"]["pages"])
        self.assertEqual(
            metadata["document"]["text_characters"], self.result["stats"]["text_characters"]
        )
        self.assertEqual([page["page"] for page in metadata["pages"]], [1, 2])
        self.assertEqual(len(metadata["images"]), self.result["stats"]["images_total"])
        self.assertGreaterEqual(len(metadata["images"]), 2)
        markdown = (output / "document.md").read_text(encoding="utf-8")
        # An omitted image is still named in an HTML comment, so inclusion has to
        # be judged from actual Markdown image references.
        references = set(re.findall(r"!\[[^\]]*\]\((images/[^)]+)\)", markdown))
        for image in metadata["images"]:
            self.assertTrue(image["path"].startswith("images/"))
            self.assertTrue((output / image["path"]).is_file())
            self.assertIn(image["confidence_band"], ("high", "medium", "low"))
            self.assertIsInstance(image["included_in_markdown"], bool)
            self.assertEqual(image["included_in_markdown"], image["path"] in references)
            if not image["included_in_markdown"]:
                self.assertEqual(image["omission_reason"], "decorative_or_logo")
        # Compact: no raw signal payloads, evidence ledgers or per-block text.
        self.assertNotIn("signals", metadata["images"][0])
        self.assertNotIn("evidence_ledger", metadata["images"][0])
        self.assertNotIn("text", metadata)

    def test_image_metadata_carries_retrieval_fields(self):
        metadata = self.read(self.default_output, "metadata.json")
        image = metadata["images"][0]
        for field in (
            "id",
            "asset_sha256",
            "page",
            "pages",
            "type",
            "confidence",
            "confidence_band",
            "caption",
            "ocr_text",
            "width",
            "height",
            "duplicate_of",
            "occurrences",
            "included_in_markdown",
            "omission_reason",
        ):
            self.assertIn(field, image)
        record = self.result["images"][0]
        self.assertEqual(image["confidence"], record["classification_confidence"])
        self.assertEqual(image["asset_sha256"], record["sha256"])
        self.assertEqual(image["type"], record["semantic_class"])

    def test_same_visual_role_gets_the_same_decision_on_every_page(self):
        """The same chart and the same logo appear on both pages.

        Deduplication is byte-exact, and Docling re-encodes an embedded image
        per page, so these occurrences are separate assets. What must still hold
        is that the parser reaches the same decision for the same content.
        """
        metadata = self.read(self.default_output, "metadata.json")
        by_type: dict[str, list[dict]] = {}
        for image in metadata["images"]:
            by_type.setdefault(image["type"], []).append(image)
        self.assertEqual(sorted(by_type), ["chart", "logo"])
        for image_type, group in by_type.items():
            self.assertEqual(len(group), 2, image_type)
            self.assertEqual(len({item["page"] for item in group}), 2, image_type)
            self.assertEqual(len({item["included_in_markdown"] for item in group}), 1, image_type)
            self.assertEqual(len({item["confidence"] for item in group}), 1, image_type)
        for image in by_type["chart"]:
            self.assertTrue(image["included_in_markdown"])
            self.assertIsNone(image["omission_reason"])
        for image in by_type["logo"]:
            self.assertFalse(image["included_in_markdown"])
        self.assertEqual(
            self.result["stats"]["images_in_md"],
            len([item for item in metadata["images"] if item["included_in_markdown"]]),
        )

    def test_decision_log_has_one_json_line_per_image(self):
        output = self.default_output
        lines = (output / "logs" / "image_decisions.jsonl").read_text(
            encoding="utf-8"
        ).splitlines()
        self.assertEqual(len(lines), self.result["stats"]["images_total"])
        for line in lines:
            entry = json.loads(line)
            for field in (
                "ref",
                "page",
                "path",
                "semantic_class",
                "confidence",
                "confidence_band",
                "method",
                "reason",
                "included_in_markdown",
                "evidence",
            ):
                self.assertIn(field, entry)
        for entry in (json.loads(line) for line in lines):
            if not entry["included_in_markdown"]:
                self.assertTrue(entry["omit_reason"])

    def test_manifest_lists_everything_and_checksums_verify(self):
        from parse_docling_slim import _sha256_file

        result, output = self.parse_with("full", emit_canonical=True)
        manifest = self.read(output, "manifest.json")
        on_disk = {
            path.relative_to(output).as_posix()
            for path in output.rglob("*")
            if path.is_file()
        }
        listed = {item["path"] for item in manifest["artifacts"]}
        self.assertEqual(listed | {"manifest.json"}, on_disk)
        for item in manifest["artifacts"]:
            self.assertEqual(_sha256_file(output / item["path"]), item["sha256"])
            self.assertEqual((output / item["path"]).stat().st_size, item["bytes"])
        self.assertEqual(manifest["integrity"]["errors"], [])
        self.assertEqual(manifest["integrity"]["unlisted_files"], [])
        self.assertEqual(manifest["status"], "complete")
        self.assertEqual(result["validation"]["status"], "pass")

    def test_validation_can_be_disabled_but_is_recorded(self):
        _, output = self.parse_with("novalidate", validate_output=False)
        self.assertEqual(self.read(output, "manifest.json")["validation"]["status"], "skipped")
        self.assertEqual(self.read(output, "metadata.json")["validation"]["status"], "skipped")

    def test_type_aware_warnings_are_bucketed(self):
        from parse_docling_slim import WARNING_WARNING, ParserConfig, _type_aware_warnings

        confident = make_record(
            ocr="14%\n19%\n35%\n28%",
            semantic_class="chart",
            classification_confidence=0.93,
            classification_confidence_band="high",
            include_in_md=True,
        )
        uncertain = make_record(
            ocr="\ufffd\ufffd 1 2\n| | \ufffd",
            width=300,
            height=200,
            semantic_class="unclassified_image",
            classification_confidence=0.0,
            classification_confidence_band="low",
            include_in_md=True,
            relative_path="images/image_0002.png",
        )
        findings = _type_aware_warnings([confident, uncertain], ParserConfig())
        kinds = {item["kind"] for item in findings}
        self.assertIn("uncertain_classification", kinds)
        warning = next(item for item in findings if item["kind"] == "uncertain_classification")
        self.assertEqual(warning["severity"], WARNING_WARNING)
        self.assertIn("image_0002.png", warning["message"])
        self.assertTrue(warning["suggestion"])
        self.assertNotIn("medium_confidence_classification", kinds)

    def test_low_confidence_dominance_is_flagged(self):
        from parse_docling_slim import WARNING_WARNING, ParserConfig, _type_aware_warnings

        records = [
            make_record(
                order=index,
                relative_path=f"images/image_{index:04d}.png",
                path=f"images/image_{index:04d}.png",
                ocr="\ufffd\ufffd 1 2",
                width=400,
                height=300,
                semantic_class="unclassified_image",
                classification_confidence=0.0,
                classification_confidence_band="low",
                include_in_md=True,
            )
            for index in range(4)
        ]
        findings = _type_aware_warnings(records, ParserConfig())
        dominance = next(
            item for item in findings if item["kind"] == "low_confidence_dominates"
        )
        self.assertEqual(dominance["severity"], WARNING_WARNING)
        self.assertIn("text layer", dominance["suggestion"])
        self.assertIn("4 of 4", dominance["message"])


class ValidationTests(unittest.TestCase):
    """I, J: broken images, missing markers. Plus semantic-loss warnings."""

    class FakeDoc:
        def __init__(self, pages):
            self.pages = {
                page: SimpleNamespace(size=SimpleNamespace(width=600, height=800), image=None)
                for page in pages
            }
            self.marker_counts = {page: 0 for page in pages}

        def export_to_markdown(self, page_no=None, **kwargs):
            if page_no is None:
                return IMAGE_MARKER * sum(self.marker_counts.values())
            return IMAGE_MARKER * self.marker_counts.get(page_no, 0)

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.output = Path(self.temp_dir.name)
        (self.output / "images").mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (10, 10), "white").save(self.output / "images" / "image_0001.png")

    def tearDown(self):
        self.temp_dir.cleanup()

    def validate(self, markdown, records, tables=None, pages=(1,)):
        # Validation checks the artifacts a consumer will read, so the test
        # writes the Markdown exactly as the parser would before validating it.
        (self.output / "document.md").write_text(markdown, encoding="utf-8")
        doc = self.FakeDoc(pages)
        for record in records:
            page = record.get("page")
            if page in doc.marker_counts:
                doc.marker_counts[page] += 1
        # These unit tests exercise the Markdown/image rules in isolation, so the
        # metadata and canonical artifacts are declared as not expected here.
        from parse_docling_slim import ParserConfig

        return _validate_outputs(
            self.output,
            markdown,
            doc,
            records,
            tables or [],
            config=ParserConfig(emit_metadata=False, emit_canonical=False),
        )

    def test_i_missing_image_asset_is_an_error(self):
        records = [make_record(relative_path="images/image_0002.png", path="images/image_0002.png")]
        result = self.validate("![a](images/image_0002.png)", records)
        self.assertEqual(result["status"], "fail")
        self.assertTrue(any("missing" in error for error in result["errors"]))

    def test_i_broken_relative_path_is_an_error(self):
        records = [make_record(include_in_md=False)]
        result = self.validate("![a](images/missing_asset.png)", records)
        self.assertEqual(result["status"], "fail")

    def test_i_non_asset_image_reference_is_an_error(self):
        result = self.validate("![a](file:///etc/passwd)", [], pages=(1,))
        self.assertEqual(result["status"], "fail")

    def test_i_base64_payload_is_an_error(self):
        markdown = "<!-- PAGE: 1 -->\n\n![a](data:image/png;base64,AAAA)\n"
        result = self.validate(markdown, [], pages=(1,))
        self.assertEqual(result["status"], "fail")

    def test_j_missing_page_marker_is_an_error(self):
        markdown = "<!-- PAGE: 1 -->\n\n<!-- PAGE: 3 -->\n"
        result = self.validate(markdown, [], pages=(1, 2, 3))
        self.assertEqual(result["status"], "fail")

    def test_j_non_sequential_page_markers_are_detected(self):
        markdown = "<!-- PAGE: 1 -->\n<!-- PAGE: 2 -->\n<!-- PAGE: 4 -->\n"
        result = self.validate(markdown, [], pages=(1, 2, 4))
        self.assertEqual(result["status"], "pass")
        markdown = "<!-- PAGE: 1 -->\n<!-- PAGE: 3 -->\n"
        result = self.validate(markdown, [], pages=(1, 2, 3))
        self.assertEqual(result["status"], "fail")
        markdown = "<!-- PAGE: 1 -->\n<!-- PAGE: 1 -->\n<!-- PAGE: 2 -->\n"
        result = self.validate(markdown, [], pages=(1, 2))
        self.assertEqual(result["status"], "fail")

    def test_j_unresolved_marker_is_an_error(self):
        result = self.validate("<!-- PAGE: 1 -->\n\n<!-- image -->\n", [], pages=(1,))
        self.assertEqual(result["status"], "fail")
        self.assertTrue(any("marker" in error for error in result["errors"]))

    def test_semantic_loss_is_reported_as_warning_not_error(self):
        record = make_record(
            page=32,
            ocr="% of physicians\n14%\n19%\n35%\n28%",
            include_in_md=False,
            omit_reason="decorative_or_logo",
            relative_path="images/image_0042.png",
            path="images/image_0042.png",
        )
        record["semantic_class"] = "chart"
        result = self.validate("<!-- PAGE: 32 -->\n", [record], pages=(32,))
        self.assertEqual(result["status"], "warning")
        self.assertEqual(result["errors"], [])
        self.assertTrue(
            any("image_0042.png" in warning and "page 32" in warning for warning in result["warnings"])
        )
        self.assertTrue(result["semantic_loss"])
        self.assertIn("quantitative OCR content (percentages)", result["semantic_loss"][0]["evidence"])

    def test_decorative_omission_is_not_a_semantic_loss(self):
        record = make_record(ocr="IQVIA", width=162, height=32, include_in_md=False)
        record["semantic_class"] = "logo"
        self.assertEqual(_semantic_loss_warnings([record], self.output), [])

    def test_included_image_without_markdown_reference_is_an_error(self):
        record = make_record(include_in_md=True)
        result = self.validate("<!-- PAGE: 1 -->\n", [record])
        self.assertEqual(result["status"], "fail")

    def test_valid_output_passes(self):
        record = make_record(include_in_md=True, ocr="14%")
        markdown = "<!-- PAGE: 1 -->\n\n![a](images/image_0001.png)\n"
        result = self.validate(markdown, [record], tables=[{"ref": "#/tables/0", "markdown": "| a |", "rows": [["1"]], "image_path": None}])
        self.assertEqual(result["status"], "pass")
        self.assertEqual(result["errors"], [])


class TablePolicyTests(unittest.TestCase):
    """E: complex tables keep structure and the original image."""

    def test_representations(self):
        self.assertEqual(_table_representation({"markdown": "| a |", "rows": [["1"]], "image_path": "images/table_0001.png"}), "structured+image")
        self.assertEqual(_table_representation({"markdown": "| a |", "rows": [["1"]], "image_path": None}), "structured")
        self.assertEqual(_table_representation({"markdown": "", "rows": [], "image_path": "images/table_0001.png"}), "image")
        self.assertEqual(_table_representation({"markdown": "", "rows": [], "image_path": None}), "none")

    def test_complex_table_without_image_is_flagged(self):
        record = make_record(ocr="14%")
        table = {
            "type": "table",
            "ref": "#/tables/0",
            "page": 1,
            "markdown": "",
            "rows": [],
            "complex": True,
            "image_path": None,
        }
        doc = ValidationTests.FakeDoc([1])
        from parse_docling_slim import ParserConfig

        output = Path(tempfile.mkdtemp())
        (output / "document.md").write_text("<!-- PAGE: 1 -->\n", encoding="utf-8")
        result = _validate_outputs(
            output,
            "<!-- PAGE: 1 -->\n",
            doc,
            [],
            [table],
            config=ParserConfig(emit_metadata=False, emit_canonical=False),
        )
        self.assertEqual(result["status"], "warning")
        self.assertTrue(any("Complex table" in warning for warning in result["warnings"]))
        self.assertEqual(table["representation"], "none")


class OfficeFixtureTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        image_path = self.root / "fixture.png"
        image = Image.new("RGB", (900, 500), "white")
        draw = ImageDraw.Draw(image)
        draw.rectangle((30, 30, 870, 470), outline="black", width=4)
        draw.text((60, 70), "Quarterly chart", fill="black")
        draw.text((100, 180), "Q1  42%", fill="black")
        image.save(image_path)
        self.image_path = image_path

    def tearDown(self):
        self.temp_dir.cleanup()

    def make_docx(self):
        from docx import Document
        from docx.shared import Inches

        document = Document()
        document.add_heading("Fixture Document", level=1)
        document.add_paragraph("Introductory paragraph with enough text for conversion.")
        document.add_paragraph("First list item", style="List Bullet")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Region"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "North"
        table.cell(1, 1).text = "42%"
        document.add_picture(str(self.image_path), width=Inches(5))
        document.add_paragraph("Closing paragraph after image.")
        path = self.root / "fixture.docx"
        document.save(path)
        return path

    def make_pptx(self):
        from pptx import Presentation
        from pptx.util import Inches

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = "Fixture Slide"
        slide.placeholders[1].text = "Subtitle and body copy"
        slide.shapes.add_picture(str(self.image_path), Inches(1), Inches(2), width=Inches(8))
        second = presentation.slides.add_slide(presentation.slide_layouts[5])
        second.shapes.title.text = "Second Slide"
        textbox = second.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(1))
        textbox.text_frame.text = "Slide table follows"
        table_shape = second.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(8), Inches(1.5))
        table = table_shape.table
        table.cell(0, 0).text = "Region"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "North"
        table.cell(1, 1).text = "42%"
        path = self.root / "fixture.pptx"
        presentation.save(path)
        return path

    def make_chart_pptx(self):
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Chart Slide"
        data = CategoryChartData()
        data.categories = ["North", "South"]
        data.add_series("Series 1", (42, 58))
        slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(1),
            Inches(2),
            Inches(8),
            Inches(5),
            data,
        )
        path = self.root / "chart.pptx"
        presentation.save(path)
        return path

    def test_docx_outputs_page_table_and_image(self):
        path = self.make_docx()
        output = self.root / "docx-output"
        result = parse_document(
            str(path), str(output), enable_vision=False, emit_canonical=True, force=True
        )
        markdown = (output / "document.md").read_text(encoding="utf-8")
        canonical = json.loads((output / "canonical.json").read_text(encoding="utf-8"))
        manifest = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(canonical["validation"]["status"], "pass")
        self.assertEqual(manifest["validation"]["status"], "pass")
        self.assertTrue(any(item["path"] == "logs/parser.log" for item in manifest["artifacts"]))
        self.assertIn("<!-- PAGE: 1 -->", markdown)
        self.assertIn("| Region", markdown)
        self.assertIn("![", markdown)
        self.assertEqual(result["stats"]["tables"], 1)

    def test_pptx_outputs_slide_markers_and_image(self):
        path = self.make_pptx()
        output = self.root / "pptx-output"
        result = parse_document(
            str(path), str(output), enable_vision=False, emit_canonical=True, force=True
        )
        markdown = (output / "document.md").read_text(encoding="utf-8")
        canonical = json.loads((output / "canonical.json").read_text(encoding="utf-8"))
        self.assertEqual(canonical["validation"]["status"], "pass")
        self.assertIn("<!-- PAGE: 1 -->", markdown)
        self.assertIn("<!-- PAGE: 2 -->", markdown)
        self.assertIn("Fixture Slide", markdown)
        self.assertIn("| Region", markdown)
        self.assertIn("![", markdown)
        self.assertEqual(result["stats"]["pages"], 2)
        self.assertEqual(result["stats"]["tables"], 1)

    def test_pptx_chart_data_is_retained_without_libreoffice(self):
        path = self.make_chart_pptx()
        output = self.root / "chart-output"
        result = parse_document(
            str(path), str(output), enable_vision=False, emit_canonical=True, force=True
        )
        canonical = json.loads((output / "canonical.json").read_text(encoding="utf-8"))
        markdown = (output / "document.md").read_text(encoding="utf-8")
        self.assertEqual(canonical["validation"]["status"], "pass")
        self.assertEqual(canonical["images"][0]["classification"], "bar_chart")
        self.assertEqual(
            canonical["images"][0]["chart_data"]["rows"],
            [["North", "42"], ["South", "58"]],
        )
        self.assertIn("| North", markdown)
        self.assertEqual(result["stats"]["pages"], 1)

    def test_docx_images_are_classified_and_referenced(self):
        path = self.make_docx()
        output = self.root / "docx-images"
        result = parse_document(
            str(path), str(output), enable_vision=False, emit_canonical=True, force=True
        )
        canonical = json.loads((output / "canonical.json").read_text(encoding="utf-8"))
        image = canonical["images"][0]
        self.assertTrue(image["include_in_md"])
        self.assertTrue(image["saved"])
        self.assertEqual(image.get("omit_reason"), None)
        self.assertIn(
            image["semantic_class"],
            {"chart", "text_image", "table_image", "unclassified_image"},
        )
        self.assertEqual(result["stats"]["images_in_md"], 1)


class SourceFingerprintTests(unittest.TestCase):
    """L: source fingerprint. Plus artifact checksums and atomic output."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        source = self.root / "fingerprint.pdf"
        self.source_bytes = b"%PDF-1.7\n% fingerprint fixture\n%%EOF\n"
        source.write_bytes(self.source_bytes)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_source_fingerprint_and_checksums(self):
        output = self.root / "out"
        with patch(
            "parse_docling_slim.build_converter",
            return_value=SimpleNamespace(convert=lambda path: SimpleNamespace(document=_FakeDocument())),
        ):
            canonical = parse_document(
                str(self.root / "fingerprint.pdf"), str(output), emit_canonical=True, force=True
            )
        manifest = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
        expected = hashlib.sha256(self.source_bytes).hexdigest()
        self.assertEqual(canonical["source"]["sha256"], expected)
        self.assertEqual(canonical["source"]["filename"], "fingerprint.pdf")
        self.assertEqual(canonical["source"]["format"], "pdf")
        self.assertEqual(canonical["source"]["size_bytes"], len(self.source_bytes))
        self.assertEqual(manifest["source"]["sha256"], expected)
        self.assertEqual(manifest["status"], "complete")
        checksums = {item["path"]: item["sha256"] for item in manifest["artifacts"]}
        self.assertIn("document.md", checksums)
        self.assertIn("canonical.json", checksums)
        self.assertIn("logs/parser.log", checksums)
        self.assertEqual(
            checksums["document.md"],
            hashlib.sha256((output / "document.md").read_bytes()).hexdigest(),
        )
        self.assertNotIn("manifest.json", checksums)

    def test_failed_run_leaves_no_output_directory(self):
        output = self.root / "never-created"
        with patch(
            "parse_docling_slim.build_converter",
            side_effect=RuntimeError("conversion exploded"),
        ):
            with self.assertRaises(RuntimeError):
                parse_document(str(self.root / "fingerprint.pdf"), str(output), force=True)
        self.assertFalse(output.exists())
        self.assertEqual(list(self.root.glob(".never-created.tmp-*")), [])

    def test_non_empty_output_requires_force(self):
        output = self.root / "occupied"
        output.mkdir()
        (output / "keep.txt").write_text("x", encoding="utf-8")
        with self.assertRaises(FileExistsError):
            parse_document(str(self.root / "fingerprint.pdf"), str(output))
        self.assertTrue((output / "keep.txt").is_file())


class _FakePage:
    def __init__(self):
        self.size = SimpleNamespace(width=600, height=800)
        self.image = None


class _FakeProvenance:
    def __init__(self, page=1, bbox=(60, 60, 300, 200)):
        self.page_no = page
        self.bbox = SimpleNamespace(l=bbox[0], t=bbox[1], r=bbox[2], b=bbox[3])
        self.layer = "body"


class _FakePictureItem:
    """A picture item with a real PIL image, captions and a chart classification."""

    def __init__(self, image, page=1, caption="Preferring specialist centres", ref="#/pictures/0"):
        self.self_ref = ref
        self.prov = [_FakeProvenance(page)]
        self.children = []
        self.captions = [_FakeTextItem(caption, page, f"{ref}/caption")]
        self.footers = []
        self.annotations = []
        self.meta = SimpleNamespace(
            classification=SimpleNamespace(predictions=[SimpleNamespace(class_name="chart")]),
            tabular_chart=None,
        )
        self._image = image

    def get_image(self, doc, *args, **kwargs):
        return self._image

    @property
    def image(self):
        return None


class _FakeTextItem:
    def __init__(self, text, page=1, ref="#/texts/0"):
        self.self_ref = ref
        self.text = text
        self.label = "text"
        self.prov = [_FakeProvenance(page)]
        self.children = []


class _FakeDocument:
    """Minimal DoclingDocument stand-in for output-contract tests."""

    def __init__(self, pages=(1,), items=()):
        self.pages = {int(page): _FakePage() for page in pages}
        self._items = list(items)
        self.texts = [item for item in self._items if isinstance(item, _FakeTextItem)]
        self.pictures = [item for item in self._items if isinstance(item, _FakePictureItem)]
        self.tables = []
        self.groups = []
        self.lists = []

    def iterate_items(self, **kwargs):
        # The real API yields (item, level) pairs.
        return iter((item, 1) for item in self._items)


def _fake_converter(document):
    """A converter stand-in that returns ``document`` for any input path."""
    return SimpleNamespace(convert=lambda path: SimpleNamespace(document=document))


class MarkdownReferenceTests(unittest.TestCase):
    """K: every Markdown image reference resolves to an existing asset."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.source = self.root / "reference.pdf"
        self.source.write_bytes(b"%PDF-1.7\n%%EOF\n")

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_references_resolve_and_assets_exist(self):
        output = self.root / "out"
        canonical_path = output / "canonical.json"
        output.mkdir(parents=True)
        (output / "images").mkdir()
        Image.new("RGB", (20, 20), "white").save(output / "images" / "image_0001.png")
        (output / "document.md").write_text(
            "<!-- PAGE: 1 -->\n\n![chart](images/image_0001.png)\n", encoding="utf-8"
        )
        _write_json(
            canonical_path,
            {
                "images": [
                    {
                        "relative_path": "images/image_0001.png",
                        "include_in_md": True,
                        "page": 1,
                    }
                ]
            },
        )
        markdown = (output / "document.md").read_text(encoding="utf-8")
        from parse_docling_slim import IMAGE_REF_RE

        references = [ref for ref in IMAGE_REF_RE.findall(markdown) if ref.startswith("images/")]
        self.assertEqual(references, ["images/image_0001.png"])
        for reference in references:
            self.assertTrue((output / reference).is_file())


def _write_json(path, value):
    path.write_text(json.dumps(value, indent=2) + "\n", encoding="utf-8")


class _BatchInputTests(unittest.TestCase):
    """Batch fan-out, directory walking and zip archives.

    ``parse_document`` keeps its single-file contract; everything asserted here
    is about which files reach it and where each result lands.
    """

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.output = self.root / "out"
        self.documents = self.root / "studies"
        self.documents.mkdir()

    def tearDown(self):
        self.temp_dir.cleanup()

    def write_document(self, relative, payload=b"%PDF-1.7\n%%EOF\n"):
        path = self.root / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(payload)
        return path

    def make_zip(self, name, members):
        archive_path = self.root / name
        with zipfile.ZipFile(archive_path, "w") as archive:
            for member, payload in members.items():
                archive.writestr(member, payload)
        return archive_path

    def parse(self, inputs, **overrides):
        """Run parse_many against a stub converter so the fan-out is what is tested."""
        values = {"force": True, "enable_vision": False}
        values.update(overrides)
        with patch(
            "parse_docling_slim.build_converter",
            return_value=_fake_converter(_FakeDocument()),
        ):
            return parse_many(inputs, str(self.output), **values)

    # --- resolution ---------------------------------------------------------

    def test_single_file_stays_flat_and_keeps_the_legacy_layout(self):
        source = self.write_document("studies/alpha.pdf")
        summary = self.parse([str(source)])
        self.assertEqual(summary["documents"], 1)
        self.assertEqual(summary["failed"], 0)
        # No intermediate directory: the historical single-file layout is preserved.
        self.assertEqual(summary["results"][0]["output"], str(self.output))
        self.assertTrue((self.output / "document.md").is_file())
        self.assertTrue((self.output / "manifest.json").is_file())
        self.assertFalse((self.output / "alpha").exists())

    def test_one_input_argument_accepts_several_files(self):
        first = self.write_document("studies/alpha.pdf")
        second = self.write_document("studies/beta.pdf")
        summary = self.parse([str(first), str(second)])
        self.assertEqual(summary["documents"], 2)
        self.assertEqual(summary["succeeded"], 2)
        for name in ("alpha", "beta"):
            self.assertTrue((self.output / name / "document.md").is_file())
        self.assertFalse((self.output / "document.md").exists())

    def test_directory_input_is_walked_recursively(self):
        self.write_document("studies/one/alpha.pdf")
        self.write_document("studies/two/nested/beta.pptx")
        self.write_document("studies/notes.txt")
        self.write_document("studies/three/.hidden/gamma.pdf")
        summary = self.parse([str(self.documents)])
        self.assertEqual(summary["documents"], 2)
        self.assertEqual(
            sorted(entry["source"] for entry in summary["results"]), ["alpha.pdf", "beta.pptx"]
        )

    def test_docx_and_pptx_are_collected_alongside_pdf(self):
        self.write_document("studies/a.docx")
        self.write_document("studies/b.pptx")
        summary = self.parse([str(self.documents)])
        self.assertEqual(summary["documents"], 2)

    def test_zip_archive_is_extracted_and_parsed(self):
        archive = self.make_zip(
            "studies.zip",
            {
                "studies/alpha.pdf": b"%PDF-1.7\nalpha\n%%EOF\n",
                "studies/beta.pdf": b"%PDF-1.7\nbeta\n%%EOF\n",
                "studies/readme.md": b"not a document",
            },
        )
        summary = self.parse([str(archive)])
        self.assertEqual(summary["documents"], 2)
        self.assertEqual(
            sorted(entry["source"] for entry in summary["results"]), ["alpha.pdf", "beta.pdf"]
        )
        self.assertTrue((self.output / "alpha" / "document.md").is_file())
        self.assertTrue((self.output / "beta" / "document.md").is_file())

    def test_zip_tolerates_os_metadata_and_nested_folders(self):
        archive = self.make_zip(
            "messy.zip",
            {
                "__MACOSX/._alpha.pdf": b"resource fork",
                "studies/.DS_Store": b"junk",
                "studies/alpha.pdf": b"%PDF-1.7\n%%EOF\n",
                "studies/deep/nested/beta.pdf": b"%PDF-1.7\n%%EOF\n",
            },
        )
        summary = self.parse([str(archive)])
        self.assertEqual(summary["documents"], 2)

    def test_zip_and_directory_can_be_mixed(self):
        self.write_document("studies/alpha.pdf")
        archive = self.make_zip("more.zip", {"beta.pdf": b"%PDF-1.7\n%%EOF\n"})
        summary = self.parse([str(self.documents), str(archive)])
        self.assertEqual(summary["documents"], 2)
        self.assertEqual(
            sorted(entry["source"] for entry in summary["results"]), ["alpha.pdf", "beta.pdf"]
        )

    def test_the_same_file_listed_twice_is_parsed_once(self):
        source = self.write_document("studies/alpha.pdf")
        summary = self.parse([str(source), str(self.documents)])
        self.assertEqual(summary["documents"], 1)

    # --- archive safety -----------------------------------------------------

    def test_path_traversal_member_is_refused(self):
        archive = self.make_zip(
            "evil.zip",
            {
                "../escaped.pdf": b"%PDF-1.7\n%%EOF\n",
                "studies/alpha.pdf": b"%PDF-1.7\n%%EOF\n",
            },
        )
        summary = self.parse([str(archive)])
        self.assertEqual(summary["documents"], 1)
        self.assertEqual(summary["results"][0]["source"], "alpha.pdf")
        self.assertFalse((self.root / "escaped.pdf").exists())

    def test_absolute_member_is_refused(self):
        archive = self.make_zip(
            "absolute.zip",
            {"/etc/escaped.pdf": b"%PDF-1.7\n%%EOF\n", "alpha.pdf": b"%PDF-1.7\n%%EOF\n"},
        )
        summary = self.parse([str(archive)])
        self.assertEqual(summary["documents"], 1)
        self.assertFalse(Path("/etc/escaped.pdf").exists())

    def test_non_document_member_is_never_written(self):
        archive = self.make_zip(
            "payload.zip",
            {
                "alpha.pdf": b"%PDF-1.7\n%%EOF\n",
                "install.sh": b"#!/bin/sh\nrm -rf /\n",
                "evil.exe": b"MZ",
            },
        )
        summary = self.parse([str(archive)])
        self.assertEqual(summary["documents"], 1)

    def test_archive_with_too_many_members_is_rejected(self):
        with patch("parse_docling_slim.MAX_ARCHIVE_MEMBERS", 2):
            archive = self.make_zip(
                "crowded.zip",
                {f"doc{index}.pdf": b"%PDF-1.7\n%%EOF\n" for index in range(5)},
            )
            with self.assertRaises(ValueError) as caught:
                self.parse([str(archive)])
        self.assertIn("above the", str(caught.exception))

    def test_oversized_archive_is_rejected(self):
        with patch("parse_docling_slim.MAX_ARCHIVE_BYTES", 16):
            archive = self.make_zip("big.zip", {"alpha.pdf": b"x" * 4096})
            with self.assertRaises(ValueError) as caught:
                self.parse([str(archive)])
        self.assertIn("expands past", str(caught.exception))

    def test_understated_member_size_cannot_defeat_the_cap(self):
        """A zip header may lie about its size, so real bytes must be capped too."""
        archive_path = self.root / "liar.zip"
        payload = b"%PDF-1.7\n" + b"x" * 5000
        with zipfile.ZipFile(archive_path, "w") as archive:
            info = zipfile.ZipInfo("alpha.pdf")
            info.file_size = 10  # deliberately understated
            archive.writestr(info, payload)
        with patch("parse_docling_slim.MAX_ARCHIVE_BYTES", 64):
            with self.assertRaises(ValueError) as caught:
                self.parse([str(archive_path)])
        self.assertIn("expands past", str(caught.exception))
        # The partial member must not be left behind in the extraction root.
        self.assertEqual(
            list(Path(tempfile.gettempdir()).glob("docling-slim-archive-*")), []
        )

    def test_temporary_extraction_is_removed_after_the_run(self):
        archive = self.make_zip("studies.zip", {"alpha.pdf": b"%PDF-1.7\n%%EOF\n"})
        before = set(Path(tempfile.gettempdir()).glob("docling-slim-archive-*"))
        self.parse([str(archive)])
        after = set(Path(tempfile.gettempdir()).glob("docling-slim-archive-*"))
        self.assertEqual(after, before)

    # --- failure handling ---------------------------------------------------

    def test_failed_document_is_recorded_and_the_batch_continues(self):
        good = self.write_document("studies/alpha.pdf")
        broken = self.write_document("studies/broken.pdf", b"not really a pdf")
        calls = []

        def flaky(_render_office_charts=False, **_kwargs):
            calls.append(1)
            if len(calls) == 1:
                return _fake_converter(_FakeDocument())
            raise RuntimeError("conversion exploded")

        with patch("parse_docling_slim.build_converter", side_effect=flaky):
            summary = parse_many(
                [str(broken), str(good)], str(self.output), force=True, enable_vision=False
            )
        self.assertEqual(summary["documents"], 2)
        self.assertEqual(summary["succeeded"], 1)
        self.assertEqual(summary["failed"], 1)
        failed = next(entry for entry in summary["results"] if entry["status"] == "error")
        self.assertIn("RuntimeError", failed["error"])
        # The successful document still produced a complete bundle.
        self.assertTrue((self.output / "alpha" / "document.md").is_file())

    def test_continue_on_error_false_stops_at_the_first_failure(self):
        self.write_document("studies/alpha.pdf")
        self.write_document("studies/beta.pdf")
        with patch(
            "parse_docling_slim.build_converter",
            side_effect=RuntimeError("conversion exploded"),
        ):
            with self.assertRaises(RuntimeError):
                parse_many(
                    [str(self.documents)],
                    str(self.output),
                    force=True,
                    enable_vision=False,
                    continue_on_error=False,
                )

    def test_colliding_names_get_distinct_directories(self):
        archive = self.make_zip(
            "clash.zip",
            {"a/study.pdf": b"%PDF-1.7\n%%EOF\n", "b/study.pdf": b"%PDF-1.7\n%%EOF\n"},
        )
        summary = self.parse([str(archive)])
        self.assertEqual(summary["documents"], 2)
        outputs = sorted(entry["output"] for entry in summary["results"])
        self.assertEqual(
            outputs, [str(self.output / "study"), str(self.output / "study_2")]
        )
        for destination in outputs:
            self.assertTrue((Path(destination) / "document.md").is_file())

    def test_awkward_names_are_sanitised_into_safe_directories(self):
        archive = self.make_zip(
            "odd.zip",
            {
                "Q1 2024: results (final)/v1.2 draft.pdf": b"%PDF\n",
                "Q1 2024: results (final)/notes.docx": b"%PDF\n",
            },
        )
        summary = self.parse([str(archive)])
        self.assertEqual(summary["documents"], 2)
        for entry in summary["results"]:
            destination = Path(entry["output"])
            self.assertEqual(destination.parent, self.output)
            self.assertNotIn(" ", destination.name)
            self.assertNotIn(":", destination.name)
            self.assertNotIn("/", destination.name)
            self.assertTrue((destination / "document.md").is_file())

    # --- contract guards ----------------------------------------------------

    def test_totals_aggregate_across_documents(self):
        self.write_document("studies/alpha.pdf")
        self.write_document("studies/beta.pdf")
        summary = self.parse([str(self.documents)])
        self.assertEqual(summary["totals"]["pages"], 2)
        self.assertEqual(summary["totals"]["vision_cost_usd"], 0.0)
        self.assertEqual(summary["output"], str(self.output))

    def test_unknown_option_is_rejected(self):
        source = self.write_document("studies/alpha.pdf")
        with self.assertRaises(TypeError) as caught:
            self.parse([str(source)], not_a_real_option=True)
        self.assertIn("not_a_real_option", str(caught.exception))

    def test_empty_selection_is_rejected(self):
        with self.assertRaises(ValueError):
            self.parse([str(self.documents)])

    def test_missing_input_is_rejected(self):
        with self.assertRaises(FileNotFoundError):
            self.parse([str(self.root / "nope.pdf")])

    def test_unsupported_single_file_is_rejected(self):
        source = self.write_document("studies/notes.txt", b"text")
        with self.assertRaises(ValueError) as caught:
            self.parse([str(source)])
        self.assertIn("Unsupported format", str(caught.exception))

    def test_empty_input_sequence_is_rejected(self):
        with self.assertRaises(ValueError):
            self.parse([])

    def test_a_bare_string_is_accepted_as_one_input(self):
        source = self.write_document("studies/alpha.pdf")
        with patch(
            "parse_docling_slim.build_converter",
            return_value=_fake_converter(_FakeDocument()),
        ):
            summary = parse_many(
                str(source), str(self.output), force=True, enable_vision=False
            )
        self.assertEqual(summary["documents"], 1)


class _CliInputTests(unittest.TestCase):
    """The CLI accepts several inputs without changing the single-file path."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_one_or_more_inputs_are_accepted(self):
        parser = _build_cli()
        self.assertEqual(parser.parse_args(["a.pdf"]).inputs, ["a.pdf"])
        self.assertEqual(parser.parse_args(["a.pdf", "b.pdf"]).inputs, ["a.pdf", "b.pdf"])
        self.assertEqual(parser.parse_args(["studies/"]).inputs, ["studies/"])
        self.assertEqual(parser.parse_args(["studies.zip"]).inputs, ["studies.zip"])

    def test_input_is_still_required(self):
        with self.assertRaises(SystemExit):
            _build_cli().parse_args([])

    def test_vision_flags_still_parse_alongside_many_inputs(self):
        args = _build_cli().parse_args(["a.pdf", "b.pdf", "--vision", "--force"])
        self.assertTrue(args.enable_vision)
        self.assertTrue(args.force)
        self.assertEqual(args.output, "output_docling")


if __name__ == "__main__":
    unittest.main()
